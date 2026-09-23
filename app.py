import os, json, hmac, base64, time, uuid
from io import BytesIO
from pathlib import Path
from typing import Any
from collections import defaultdict, deque

import httpx
from fastapi import FastAPI, HTTPException, Header, UploadFile, File, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"

app = FastAPI(title="NovaMind AI", version="5.0")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

AI_BASE_URL = os.getenv("AI_BASE_URL", "").rstrip("/")
AI_CHAT_URL = os.getenv("AI_CHAT_URL", "").strip()
AI_API_KEY = os.getenv("AI_API_KEY", "").strip()
AI_MODEL = os.getenv("AI_MODEL", "openai/gpt-oss-20b").strip()
AI_VISION_MODEL = os.getenv("AI_VISION_MODEL", "qwen/qwen3.8-27b").strip()
ACCESS_CODE = os.getenv("ACCESS_CODE", "").strip()
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "15"))

ATTACHMENTS: dict[str, dict[str, Any]] = {}
RATE: dict[str, deque] = defaultdict(deque)

TEXT_EXTS = {".txt",".md",".csv",".json",".html",".htm",".py",".js",".ts",".tsx",".jsx",".css",".java",".c",".cpp",".h",".hpp",".rs",".go",".sql",".xml",".yaml",".yml"}
IMAGE_EXTS = {".png",".jpg",".jpeg",".webp",".gif"}

class Message(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    messages: list[Message] = Field(min_length=1, max_length=80)
    mode: str = "auto"
    web: bool = False
    code: bool = False
    attachment_ids: list[str] = Field(default_factory=list, max_length=6)

def authorize(code: str | None):
    if ACCESS_CODE and (not code or not hmac.compare_digest(code, ACCESS_CODE)):
        raise HTTPException(status_code=401, detail="Código de acesso inválido.")

def limit(request: Request, bucket: str, maximum: int, seconds: int):
    ip = request.client.host if request.client else "unknown"
    key = f"{bucket}:{ip}"
    q = RATE[key]
    now = time.time()
    while q and q[0] < now - seconds:
        q.popleft()
    if len(q) >= maximum:
        raise HTTPException(status_code=429, detail="Muitas solicitações. Aguarde um pouco e tente novamente.")
    q.append(now)

def provider_url() -> str:
    if AI_CHAT_URL:
        return AI_CHAT_URL
    if AI_BASE_URL:
        return AI_BASE_URL + "/chat/completions"
    return ""

def is_groq() -> bool:
    return "api.groq.com" in (AI_BASE_URL + AI_CHAT_URL)

def clean_attachments():
    cutoff = time.time() - 6 * 3600
    stale = [k for k,v in ATTACHMENTS.items() if v.get("created",0) < cutoff]
    for k in stale:
        ATTACHMENTS.pop(k, None)

def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext in TEXT_EXTS:
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        reader = PdfReader(BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        doc = Document(BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".xlsx":
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"## Planilha: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                out.append("\t".join(vals))
                if sum(len(x) for x in out) > 220000:
                    break
            if sum(len(x) for x in out) > 220000:
                break
        return "\n".join(out)
    raise ValueError("Formato não suportado para extração de texto.")

@app.get("/")
async def root():
    return FileResponse(STATIC_DIR / "index.html")

@app.get("/healthz")
async def healthz():
    return {"ok": True, "version": "5.0"}

@app.get("/api/status")
async def status():
    configured = bool(provider_url() and AI_API_KEY and AI_MODEL)
    return {
        "online": True,
        "ai_configured": configured,
        "web_configured": configured and is_groq(),
        "code_configured": configured and is_groq(),
        "vision_configured": bool(provider_url() and AI_API_KEY and AI_VISION_MODEL),
        "uploads_configured": True,
        "protected": bool(ACCESS_CODE),
        "model": AI_MODEL if configured else None,
        "vision_model": AI_VISION_MODEL if configured else None,
    }

@app.post("/api/upload")
async def upload(
    request: Request,
    file: UploadFile = File(...),
    x_access_code: str | None = Header(default=None)
):
    authorize(x_access_code)
    limit(request, "upload", 20, 600)
    clean_attachments()
    data = await file.read()
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"Arquivo maior que {MAX_UPLOAD_MB} MB.")
    name = file.filename or "arquivo"
    ext = Path(name).suffix.lower()
    aid = uuid.uuid4().hex
    if ext in IMAGE_EXTS:
        mime = file.content_type or ("image/jpeg" if ext in {".jpg",".jpeg"} else f"image/{ext.lstrip('.')}")
        ATTACHMENTS[aid] = {
            "created": time.time(), "name": name, "kind": "image",
            "data_url": "data:" + mime + ";base64," + base64.b64encode(data).decode("ascii")
        }
        return {"id": aid, "name": name, "kind": "image", "size": len(data)}
    try:
        text = extract_text(name, data)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Formato não suportado ou arquivo inválido.") from exc
    text = text[:220000]
    ATTACHMENTS[aid] = {"created": time.time(), "name": name, "kind": "document", "text": text}
    return {"id": aid, "name": name, "kind": "document", "characters": len(text), "size": len(data)}

@app.delete("/api/upload/{attachment_id}")
async def delete_upload(attachment_id: str, x_access_code: str | None = Header(default=None)):
    authorize(x_access_code)
    ATTACHMENTS.pop(attachment_id, None)
    return {"ok": True}

@app.post("/api/chat")
async def chat(
    request: Request,
    req: ChatRequest,
    x_access_code: str | None = Header(default=None)
):
    authorize(x_access_code)
    limit(request, "chat", 30, 600)
    clean_attachments()
    url = provider_url()
    if not (url and AI_API_KEY and AI_MODEL):
        raise HTTPException(
            status_code=503,
            detail="A NovaMind está publicada, mas falta adicionar AI_API_KEY nas Variables do Railway."
        )

    messages: list[dict[str, Any]] = [m.model_dump() for m in req.messages]
    docs = []
    images = []
    for aid in req.attachment_ids:
        item = ATTACHMENTS.get(aid)
        if not item:
            continue
        if item["kind"] == "document":
            docs.append(item)
        elif item["kind"] == "image":
            images.append(item)

    mode_hint = {
        "fast": "Seja rápido e objetivo. Não omita informações essenciais.",
        "deep": "Analise com profundidade. Verifique cálculos e inconsistências. Mostre apenas passos verificáveis úteis, nunca pensamento interno privado.",
        "auto": "Adapte automaticamente a profundidade e o nível de detalhe à tarefa.",
    }.get(req.mode, "Adapte automaticamente a profundidade à tarefa.")

    system = (
        "Você é NovaMind, uma IA pessoal multimodal, clara, útil e cuidadosa. "
        "Responda em português do Brasil por padrão. "
        "Quando não souber, diga que não sabe. Não invente fontes, resultados de ferramentas ou conteúdo de arquivos. "
        + mode_hint
    )

    if docs:
        budget = 50000
        blocks = []
        used = 0
        for d in docs:
            chunk = d["text"][:max(0, budget-used)]
            if not chunk:
                break
            blocks.append(f"\n--- ARQUIVO: {d['name']} ---\n{chunk}")
            used += len(chunk)
        system += "\n\nUse estes arquivos como contexto quando forem relevantes:" + "".join(blocks)

    messages.insert(0, {"role": "system", "content": system})

    model = AI_MODEL
    if images:
        model = AI_VISION_MODEL
        last_user = None
        for i in range(len(messages)-1, -1, -1):
            if messages[i].get("role") == "user":
                last_user = messages[i]
                break
        if last_user:
            text_content = last_user.get("content","")
            multimodal = [{"type":"text","text": text_content}]
            for img in images[:3]:
                multimodal.append({"type":"image_url","image_url":{"url":img["data_url"]}})
            last_user["content"] = multimodal

    reasoning = {"fast":"low","auto":"medium","deep":"high"}.get(req.mode, "medium")
    payload: dict[str, Any] = {
        "model": model,
        "messages": messages,
        "reasoning_effort": reasoning,
    }

    tools = []
    if not images and is_groq():
        if req.web:
            tools.append({"type":"browser_search"})
        if req.code:
            tools.append({"type":"code_interpreter"})
    if tools:
        payload["tools"] = tools
        payload["tool_choice"] = "required" if len(tools) == 1 else "auto"

    headers = {
        "Authorization": f"Bearer {AI_API_KEY}",
        "Content-Type": "application/json",
    }

    async def generate_nonstream():
        try:
            body = dict(payload)
            body["stream"] = False
            async with httpx.AsyncClient(timeout=httpx.Timeout(120.0, connect=20.0)) as client:
                r = await client.post(url, headers=headers, json=body)
            if r.status_code >= 400:
                detail = "Falha no provedor de IA."
                try:
                    info = r.json()
                    detail = info.get("error",{}).get("message") or info.get("message") or detail
                except Exception:
                    pass
                yield json.dumps({"type":"error","message":detail[:500]}, ensure_ascii=False) + "\n"
                return
            obj = r.json()
            content = obj.get("choices",[{}])[0].get("message",{}).get("content") or ""
            step = 180
            for i in range(0, len(content), step):
                yield json.dumps({"type":"delta","text":content[i:i+step]}, ensure_ascii=False) + "\n"
            yield json.dumps({"type":"done","model":model,"tools":[t["type"] for t in tools]}, ensure_ascii=False) + "\n"
        except Exception:
            yield json.dumps({"type":"error","message":"Falha ao conectar ao provedor de IA."}, ensure_ascii=False) + "\n"

    async def generate_stream():
        try:
            body = dict(payload)
            body["stream"] = True
            async with httpx.AsyncClient(timeout=httpx.Timeout(120.0, connect=20.0)) as client:
                async with client.stream("POST", url, headers=headers, json=body) as r:
                    if r.status_code >= 400:
                        raw = await r.aread()
                        msg = f"Provedor respondeu HTTP {r.status_code}."
                        try:
                            info = json.loads(raw.decode())
                            msg = info.get("error",{}).get("message") or info.get("message") or msg
                        except Exception:
                            pass
                        yield json.dumps({"type":"error","message":msg[:500]}, ensure_ascii=False) + "\n"
                        return
                    async for line in r.aiter_lines():
                        if not line or not line.startswith("data:"):
                            continue
                        data = line[5:].strip()
                        if data == "[DONE]":
                            break
                        try:
                            obj = json.loads(data)
                            delta = obj.get("choices",[{}])[0].get("delta",{}).get("content")
                            if delta:
                                yield json.dumps({"type":"delta","text":delta}, ensure_ascii=False) + "\n"
                        except Exception:
                            continue
            yield json.dumps({"type":"done","model":model,"tools":[]}, ensure_ascii=False) + "\n"
        except Exception:
            yield json.dumps({"type":"error","message":"Falha ao conectar ao provedor de IA."}, ensure_ascii=False) + "\n"

    generator = generate_nonstream() if (tools or images) else generate_stream()
    return StreamingResponse(generator, media_type="application/x-ndjson")
